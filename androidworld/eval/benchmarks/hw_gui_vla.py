"""Frozen task loader for the HW-GUI-VLA physical-device benchmark."""

from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


EXCLUDED_APPS = frozenset({"中国联通", "咸鱼之王"})
APP_PACKAGES = {
    "腾讯视频": "com.tencent.qqlive",
    "优酷视频": "com.youku.phone",
    "酷狗音乐": "com.kugou.android",
    "爱奇艺": "com.qiyi.video",
    "QQ音乐": "com.tencent.qqmusic",
    "红果": "com.phoenix.read",
    "芒果TV": "com.hunantv.imgo.activity",
    "网易云音乐": "com.netease.cloudmusic",
    "今日头条": "com.ss.android.article.news",
    "高德": "com.autonavi.minimap",
    "应用市场": "com.huawei.appmarket",
}


@dataclass(frozen=True)
class TaskRecord:
    key: str
    sheet: str
    excel_row: int
    case_id: str
    app: str
    instruction: str
    expected_steps: int
    max_steps: int
    count_in_primary_metric: bool
    fields: dict[str, Any]


def _slug(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-")
    return cleaned[:80] or "task"


def _json_value(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def load_tasks(
    workbook_path: Path,
    excluded_apps: set[str] | frozenset[str] = EXCLUDED_APPS,
) -> list[TaskRecord]:
    """Load every populated task except the explicitly excluded apps."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    tasks: list[TaskRecord] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or "用例编号" not in header or "任务" not in header:
            continue
        columns = {str(name): index for index, name in enumerate(header) if name}
        required = {"用例编号", "任务", "涉及APP", "预期步数", "是否统计"}
        missing = required - columns.keys()
        if missing:
            raise ValueError(
                f"Worksheet {worksheet.title!r} is missing columns: {sorted(missing)}"
            )
        for excel_row, row in enumerate(rows, start=2):
            case_id = row[columns["用例编号"]]
            instruction = row[columns["任务"]]
            if not case_id or not instruction:
                continue
            app = str(row[columns["涉及APP"]]).strip()
            if app in excluded_apps:
                continue
            if app not in APP_PACKAGES:
                raise ValueError(f"No package mapping for app {app!r}.")
            expected = row[columns["预期步数"]]
            if not isinstance(expected, (int, float)) or expected <= 0:
                raise ValueError(
                    f"Invalid expected steps at {worksheet.title}!{excel_row}: "
                    f"{expected!r}"
                )
            expected_steps = int(expected)
            fields = {
                str(name): _json_value(row[index]) for name, index in columns.items()
            }
            tasks.append(
                TaskRecord(
                    key=(
                        f"s{sheet_index + 1}-r{excel_row:04d}-"
                        f"{_slug(str(case_id))}"
                    ),
                    sheet=worksheet.title,
                    excel_row=excel_row,
                    case_id=str(case_id),
                    app=app,
                    instruction=str(instruction),
                    expected_steps=expected_steps,
                    max_steps=math.ceil(expected_steps * 1.25),
                    count_in_primary_metric=(
                        str(row[columns["是否统计"]] or "").strip().upper() != "N"
                    ),
                    fields=fields,
                )
            )
    return tasks
